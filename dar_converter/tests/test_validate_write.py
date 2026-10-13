from datetime import datetime, timedelta

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.formula import ArrayFormula

from dar.build import Builder
from dar.extract import extract_all
from dar.validate import validate_day
from dar.write import load_template, write_dar, write_review
from tests.conftest import AVI, DAY, NOA, msg


def make_template(path):
    """A tiny stand-in for the real DAR template: same layout, one day sheet."""
    wb = openpyxl.Workbook()
    dt = wb.active
    dt.title = "Daily Totals"
    dt["B2"] = datetime(2026, 8, 1)
    ws = wb.create_sheet("0805")
    ws["A2"] = "='Daily Totals'!B6"
    ws["A8"], ws["C8"] = "Shift", "Action"
    ws["H9"] = ArrayFormula("H9:H40", "=IF(LEN(C9:C40), MOD(G9:G40-F9:F40,1),)")
    for r in range(10, 41):
        for c in "DEFG":
            ws[f"{c}{r}"].number_format = "[h]:mm:ss"
    ws["A10"], ws["B10"], ws["C10"] = "Shift 3", "Old Name", "Patrol: Interior"   # manual data
    ws["D10"] = timedelta(hours=7)
    ws["H10"] = timedelta(minutes=3)                                              # cached result
    dv = DataValidation(type="list", formula1='"Shift 1,Shift 2,Shift 3"')
    dv.add("A10:A40")
    ws.add_data_validation(dv)
    wb.create_sheet("Legend")["A2"] = "Administrative Note"
    wb.save(path)


def day_result(ref, settings, messages):
    res = Builder(ref, settings).build_day(DAY, extract_all(messages, ref))
    validate_day(res, settings, ref)
    return res


def test_validation_flags_gaps_and_long_patrol(ref, settings):
    res = day_result(ref, settings, [msg("07:46", NOA, "39-42 אבי סרק נועה בדקה c2"),
                                     msg("12:10", NOA, "05-08 אבי סרק נועה בדקה c1"),
                                     msg("14:40", NOA, "10-35 אבי סרק נועה בדקה c2")])
    kinds = [i.kind for i in res.issues]
    assert "gap" in kinds
    long = [r for r in res.rows if r.start == 14 * 60 + 10 and r.action == "Patrol: Interior"]
    assert long and any("משך חריג" in f for f in long[0].flags)


def test_write_keeps_template_and_fills_rows(tmp_path, ref, settings):
    tpl = tmp_path / "tpl.xlsx"
    make_template(tpl)
    s = settings.with_overrides({"template": {"last_data_row": 40}})
    res = day_result(ref, s, [msg("07:46", NOA, "39-42 אבי סרק נועה בדקה c2"),
                              msg("11:46", NOA, "44 בדיקה לרון c2"),
                              msg("15:49", NOA, "49 תחילת ליווי צביקה במשרד")])  # unknown name -> review
    wb = load_template(tpl, s)
    where = write_dar(wb, [res], s)
    out = tmp_path / "out.xlsx"
    wb.save(out)
    write_review(tmp_path / "review.xlsx", [res], where, s, [], [])

    ws = openpyxl.load_workbook(out)["0805"]
    assert isinstance(ws["H9"].value, ArrayFormula) and ws["H9"].value.ref == "H9:H40"
    assert ws["H10"].value is None                         # stale cached value removed
    assert ws.data_validations.dataValidation               # dropdown kept
    rows = [[c.value for c in r] for r in ws.iter_rows(min_row=10, max_row=20, max_col=11)]
    rows = [r for r in rows if any(r)]
    actions = [r[2] for r in rows]
    assert actions.count("Assumed Post") == 3 and "Old Name" not in [r[1] for r in rows]
    patrol = next(r for r in rows if r[2] == "Patrol: Interior")
    assert patrol[3] == timedelta(hours=7, minutes=38) and patrol[6] == timedelta(hours=7, minutes=42)
    assert ws["D10"].number_format == "[h]:mm:ss"
    assert ws["C4"].value == 48 and ws["C5"].value == 0
    esc = next(i for i, r in enumerate(rows) if r[2] == "Escort: Foot")   # unknown name
    assert rows[esc][10].startswith("⚠ דרוש תיקון")           # visible marker, red row, note
    assert ws.cell(10 + esc, 2).fill.fgColor.rgb.endswith("FFC7CE")
    assert ws.cell(10 + esc, 3).comment is not None
    review = openpyxl.load_workbook(tmp_path / "review.xlsx")["שורות לבדיקה"]
    assert any(row[3] == "דרוש תיקון" for row in review.iter_rows(min_row=2, values_only=True))


def test_retarget_month_renames_tabs_formulas_and_dates(tmp_path):
    from dar.write import retarget_month, template_month
    wb = openpyxl.Workbook()
    dt = wb.active
    dt.title = "Daily Totals"
    for d in range(1, 32):
        dt.cell(d + 1, 2).value = datetime(2026, 8, d)
        dt.cell(d + 1, 3).value = "=TOTALS!$F$4"
    tot = wb.create_sheet("TOTALS")
    for d in (1, 30, 31):
        wb.create_sheet(f"08{d:02d}")["A2"] = f"='Daily Totals'!B{d + 1}"
    tot["C4"] = "=COUNTIFS('0801'!$A$10:$A150,\"Shift 1\")"
    tot["D4"] = "=COUNTIFS('0831'!$A$10:$A150,\"Shift 1\")"
    retarget_month(wb, 2026, 9)                       # September has 30 days
    assert [n for n in wb.sheetnames if n[:2] == "09"] == ["0901", "0930"]
    assert "0831" not in wb.sheetnames and tot["D4"].value is None
    assert tot["C4"].value == "=COUNTIFS('0901'!$A$10:$A150,\"Shift 1\")"
    assert template_month(wb) == (2026, 9) and dt["B31"].value == datetime(2026, 9, 30)
    assert dt["B32"].value is None and dt["C32"].value is None
