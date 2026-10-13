import json
from datetime import timedelta

import openpyxl
import pytest
import yaml

from dar.build import Builder
from dar.corrections import Corrections, event_problems
from dar.evaluate import compare_day, load_actual, norm_comment
from dar.extract import extract_all
from dar.normalize import Reference
from dar.pipeline import run
from tests.conftest import DAY, NOA, ROOT, msg


def test_comment_normalisation():
    assert norm_comment("Metal screening C2 :  Ron Shani - Declined to sit ") == \
        norm_comment("Metal screening c2: Ron Shani- Declined to sit")
    assert norm_comment("Escort office : Shani, Ron") == norm_comment("Escort Office: Ron Shani")


def test_actual_loader_and_field_diffs(tmp_path, ref, settings):
    wb = openpyxl.Workbook()
    wb.active.title = "Daily Totals"
    ws = wb.create_sheet("0805")
    rows = [("Shift 3", "Noa Bar", "Material Search", 7, 38, "Metal screening C2 : Avi Levi - Declined to sit"),
            ("Shift 2", "Avi Levi", "Patrol: Interior", 7, 38, "Routine Patrol"),
            ("Shift 1", "Someone", "Alarm: Forced Level 4-6", 7, 20, "rack door")]
    for i, (sh, off, act, h, m, k) in enumerate(rows, start=10):
        ws[f"A{i}"], ws[f"B{i}"], ws[f"C{i}"], ws[f"K{i}"] = sh, off, act, k
        ws[f"F{i}"] = timedelta(hours=h, minutes=m + 1)
        ws[f"E{i}"] = f"=F{i}-(1/1440)"                     # formula form used in some sheets
        ws[f"G{i}"] = timedelta(hours=h, minutes=m + 4)
    path = tmp_path / "actual.xlsx"
    wb.save(path)
    actual = load_actual(path, [DAY])
    got, oos = actual[DAY]
    assert len(got) == 2 and oos == 1
    assert got[0].times[1] == (7 * 60 + 38) * 60          # E evaluated from the formula

    res = Builder(ref, settings).build_day(DAY, extract_all([msg("07:46", NOA, "39-42 אבי סרק נועה בדקה c2")], ref))
    ev = compare_day(DAY, [r for r in res.rows if r.action != "Assumed Post"], got, oos)
    assert len(ev.pairs) == 2 and not ev.missing and not ev.extra
    # the sheet has no D, E is F-60s there (ours F-30s) and only 2 officers -> Shift 1/2 for us
    assert all(set(p.diffs) <= {"D", "E", "shift"} for p in ev.pairs)
    assert all("E" in p.diffs and "officer" not in p.diffs for p in ev.pairs)


def test_corrections_learning_and_overrides(tmp_path, ref, settings):
    m1 = msg("18:46", NOA, "42-25 אבי סרק נועה בדקה c2")
    m2 = msg("12:00", NOA, "צביקה 1153")
    m3 = msg("13:00", NOA, "שטויות")
    corr_path = tmp_path / "corrections.yaml"
    corr_path.write_text(yaml.safe_dump({
        "aliases": {"people": {"רוני": "Ron Shani"}},
        "messages": {m1.id: {"text": "42-45 אבי סרק נועה בדקה c2"},
                     m3.id: {"ignore": True},
                     m2.id: {"events": [{"kind": "check", "start": "11:52", "end": "11:53", "zone": "C1",
                                         "actor": "Noa Bar", "subjects": ["Tal Peretz"]}]}},
    }, allow_unicode=True), encoding="utf-8")
    r = Reference.from_dict(ref.to_dict())
    ref_path = tmp_path / "ref.yaml"
    r.save(ref_path)
    rr = run([m1, m2, m3], DAY, DAY, settings, r, Corrections.load(corr_path), ref_path,
             tmp_path / "examples.jsonl")
    assert rr.learned and Reference.load(ref_path).lookup("לרוני").name == "Ron Shani"
    rows = rr.results[0].rows
    assert any(r.action == "Patrol: Interior" and r.end == 18 * 60 + 45 for r in rows)
    assert any(r.comment == "Metal screening C1 : Tal Peretz - Declined to sit" for r in rows)
    examples = [json.loads(x) for x in (tmp_path / "examples.jsonl").read_text(encoding="utf-8").splitlines()]
    assert {e["id"] for e in examples} == {m1.id, m2.id}


def test_explicit_event_schema(ref):
    assert event_problems({"kind": "scan", "start": "7:39", "end": "07:42", "actor": "Noa Bar"}, ref) == []
    errs = event_problems({"kind": "walk", "start": "739", "actor": "Nobody", "color": 1}, ref)
    assert len(errs) == 4


def test_regression_examples_still_parse():
    """Every corrected message that was learned keeps producing the same events."""
    path = ROOT.parent / "data" / "examples.jsonl"
    if not path.exists():
        pytest.skip("no data/examples.jsonl yet")
    from dar.corrections import summarize
    from dar.extract import Extractor
    from dar.models import Message
    from datetime import datetime
    ref_real = Reference.load(ROOT.parent / "data" / "reference.yaml")
    ex = Extractor(ref_real)
    for line in path.read_text(encoding="utf-8").splitlines():
        item = json.loads(line)
        if "events" in item["fix"]:
            continue
        m = Message(item["id"], datetime.fromisoformat(item["ts"]), item["sender"], item["text"])
        got = [summarize(e) for e in ex.extract(m)]
        assert got == item["expected"], item["id"]
