"""Write rows into the exact DAR template (openpyxl keeps the formatting) + a separate review file."""
from __future__ import annotations

import json
import re
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .build import DayResult, sources_text
from .models import Event, Row, fmt_minutes, fmt_seconds
from .settings import Settings

DAY_SHEET = re.compile(r"^\d{4}$")
HEADER_CELLS = ("C4", "C5", "E4", "E5", "E6", "G4", "G5")
FILL_REVIEW = PatternFill("solid", fgColor="F8D7DA")     # not written to the DAR
FILL_FLAG = PatternFill("solid", fgColor="FFF3CD")       # written, please check
FILL_HEAD = PatternFill("solid", fgColor="D9E2F3")


# ------------------------------------------------------------------ template
def template_month(wb) -> tuple[int, int]:
    v = wb["Daily Totals"]["B2"].value if "Daily Totals" in wb.sheetnames else None
    if isinstance(v, datetime):
        return v.year, v.month
    first = next(n for n in wb.sheetnames if DAY_SHEET.match(n))
    return datetime.now().year, int(first[:2])


def day_sheet_name(d: date) -> str:
    return f"{d.month:02d}{d.day:02d}"


def blank_template(wb, first_row: int = 10, last_row: int = 2200) -> None:
    """Remove manual data from all day sheets but keep formats, dropdowns and formulas."""
    for ws in wb.worksheets:
        if not DAY_SHEET.match(ws.title):
            continue
        for row in ws.iter_rows(min_row=first_row, max_row=last_row, min_col=1, max_col=12):
            for c in row:
                if c.value is not None:
                    c.value = None          # H:J below row 9 only hold cached array results
                if c.comment is not None:
                    c.comment = None
        for coord in HEADER_CELLS:
            ws[coord].value = None


def load_template(path: str | Path, s: Settings, blank: bool = True):
    wb = load_workbook(path)
    if blank:
        blank_template(wb, int(s.get("template.first_data_row", 10)),
                       int(s.get("template.last_data_row", 2200)))
    return wb


def main_rows(res: DayResult, s: Settings) -> list[Row]:
    th = float(s.get("review.min_confidence_main", 0.6))
    return [r for r in res.rows if r.confidence >= th]


def write_dar(wb, results: list[DayResult], s: Settings) -> dict[tuple[date, int], int]:
    """Fill the day sheets. Returns {(day, row.seq): sheet row number}."""
    first = int(s.get("template.first_data_row", 10))
    last = int(s.get("template.last_data_row", 2200))
    notes = bool(s.get("review.cell_notes", True))
    badges = s.get("badges", {}) or {}
    where: dict[tuple[date, int], int] = {}
    for res in results:
        name = day_sheet_name(res.day)
        if name not in wb.sheetnames:
            raise ValueError(f"בתבנית אין לשונית {name} - צריך את התבנית של החודש הזה")
        ws = wb[name]
        rows = main_rows(res, s)
        if len(rows) > last - first + 1:
            raise ValueError(f"{name}: יותר שורות ({len(rows)}) ממה שהתבנית מאפשרת")
        i = first
        for r in rows:
            ws.cell(i, 1).value = r.shift
            ws.cell(i, 2).value = r.officer
            ws.cell(i, 3).value = r.action
            for col, sec in zip((4, 5, 6, 7), r.times_seconds()):
                ws.cell(i, col).value = None if sec is None else timedelta(seconds=sec)
            ws.cell(i, 11).value = r.comment
            if notes and r.flags:
                text = "לבדיקה:\n" + "\n".join(f"- {f}" for f in r.flags)
                ws.cell(i, 3).comment = Comment(text, "DAR converter", width=320, height=140)
            where[(res.day, r.seq)] = i
            i += 1
        if any(r.action != "Assumed Post" for r in rows):
            if badges.get("employee") is not None:
                ws["C4"].value = badges["employee"]
            if badges.get("visitor") is not None:
                ws["C5"].value = badges["visitor"]
    return where


# ------------------------------------------------------------------ review file
def _sheet(wb, title: str, headers: list[str], widths: list[int]):
    ws = wb.create_sheet(title)
    ws.sheet_view.rightToLeft = True
    ws.append(headers)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(1, i)
        c.font = Font(bold=True)
        c.fill = FILL_HEAD
        c.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "A2"
    return ws


def _wrap(ws):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    if ws.max_row > 1:
        ws.auto_filter.ref = ws.dimensions


def write_review(path: str | Path, results: list[DayResult], where: dict, s: Settings,
                 pending: list[Event], chatter: list[Event]) -> None:
    wb = Workbook()
    summary = wb.active
    summary.title = "סיכום"
    summary.sheet_view.rightToLeft = True
    th = float(s.get("review.min_confidence_main", 0.6))

    rv = _sheet(wb, "שורות לבדיקה",
                ["תאריך", "לשונית", "שורה ב-DAR", "ביטחון", "סיבות", "Shift", "Officer", "Action",
                 "Time Started", "Time Ended", "Comments/Notes", "הודעות מקור"],
                [11, 8, 10, 8, 45, 9, 18, 18, 11, 11, 42, 70])
    for res in results:
        for r in res.rows:
            if not r.flags and r.confidence >= th:
                continue
            row_no = where.get((res.day, r.seq))
            _, _, f, g = r.times_seconds()
            rv.append([res.day.strftime("%d/%m/%Y"), day_sheet_name(res.day),
                       row_no if row_no else "לא נכתבה", round(r.confidence, 2),
                       "\n".join(r.flags), r.shift, r.officer, r.action, fmt_seconds(f),
                       fmt_seconds(g), r.comment, sources_text(r.sources)])
            fill = FILL_FLAG if row_no else FILL_REVIEW
            for c in rv[rv.max_row]:
                c.fill = fill
    _wrap(rv)

    iss = _sheet(wb, "בעיות ביום", ["תאריך", "סוג", "תיאור", "הודעות מקור"], [11, 18, 60, 80])
    for res in results:
        for i in res.issues:
            iss.append([res.day.strftime("%d/%m/%Y"), i.kind, i.message, sources_text(i.sources)])
    _wrap(iss)

    un = _sheet(wb, "הודעות שלא זוהו",
                ["תאריך", "שעה", "שולח", "שורה", "סיבה", "ההודעה המלאה", "מזהה הודעה"],
                [11, 7, 24, 40, 30, 60, 22])
    for e in pending:
        un.append([e.msg.ts.strftime("%d/%m/%Y"), e.msg.ts.strftime("%H:%M"), e.msg.sender,
                   e.line, "\n".join(e.notes), e.msg.text, e.msg.id])
    _wrap(un)

    ros = _sheet(wb, "מאבטחים ביום", ["תאריך", "Shift", "Officer", "התחלה", "סיום", "הוסק", "הערה"],
                 [11, 9, 20, 8, 8, 6, 40])
    for res in results:
        for r in sorted(res.roster, key=lambda r: (r.slot, r.start)):
            ros.append([res.day.strftime("%d/%m/%Y"), f"Shift {r.slot}", r.officer,
                        fmt_minutes(r.start), fmt_minutes(r.end), "כן" if r.inferred else "", r.note])
    _wrap(ros)

    ch = _sheet(wb, "הודעות שאינן דיווח", ["תאריך", "שעה", "שולח", "טקסט"], [11, 7, 24, 80])
    for e in chatter:
        ch.append([e.msg.ts.strftime("%d/%m/%Y"), e.msg.ts.strftime("%H:%M"), e.msg.sender, e.line])
    _wrap(ch)

    # summary
    summary.append(["תאריך", "שורות שנכתבו", "שורות עם סימון", "שורות לבדיקה בלבד", "בעיות ביום"])
    for c in summary[1]:
        c.font = Font(bold=True)
        c.fill = FILL_HEAD
    for res in results:
        written = [r for r in res.rows if (res.day, r.seq) in where]
        flagged = [r for r in written if r.flags]
        review_only = [r for r in res.rows if (res.day, r.seq) not in where]
        summary.append([res.day.strftime("%d/%m/%Y"), len(written), len(flagged),
                        len(review_only), len(res.issues)])
    summary.append([])
    summary.append(["צהוב = נכתב לקובץ אבל כדאי לבדוק;  אדום = לא נכתב לקובץ (ביטחון נמוך)."])
    summary.append(["תיקונים: data/corrections.yaml (הסבר ב-README). הרצה חוזרת מחילה אותם."])
    for col, w in zip("ABCDE", (12, 14, 16, 18, 12)):
        summary.column_dimensions[col].width = w
    wb.save(path)


def write_pending_json(path: str | Path, pending: list[Event], msgs_by_day: dict) -> None:
    """Unrecognised lines + context, for resolving in a Claude session (no API needed)."""
    items = []
    for e in pending:
        day = e.msg.ts.date()
        ctx = [m for m in msgs_by_day.get(day, [])
               if m.id != e.msg.id and abs((m.ts - e.msg.ts).total_seconds()) <= 3600
               and (m.sender == e.msg.sender or abs((m.ts - e.msg.ts).total_seconds()) <= 1200)]
        items.append({
            "id": e.msg.id, "date": day.isoformat(), "time": e.msg.ts.strftime("%H:%M"),
            "sender": e.msg.sender, "line": e.line, "message": e.msg.text, "reasons": e.notes,
            "context": [f"[{m.ts:%H:%M}] {m.sender}: {m.text}" for m in ctx],
        })
    with open(path, "w", encoding="utf-8") as fh:
        json.dump(items, fh, ensure_ascii=False, indent=2)
