"""Compare converter output with the real (manually entered) DAR, row by row and field by field."""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path

import openpyxl

from .models import OUR_ACTIONS, Row, fmt_seconds

FIELDS = ("shift", "officer", "action", "D", "E", "F", "G", "comment")


@dataclass
class ActualRow:
    day: date
    sheet_row: int
    shift: str | None
    officer: str | None
    action: str | None
    times: tuple[int | None, int | None, int | None, int | None]   # D E F G seconds
    comment: str | None

    def value(self, f: str):
        return _field(self, f)


@dataclass
class Pair:
    pred: Row
    actual: ActualRow
    diffs: dict[str, tuple] = field(default_factory=dict)


@dataclass
class DayEval:
    day: date
    pairs: list[Pair] = field(default_factory=list)
    missing: list[ActualRow] = field(default_factory=list)      # in Excel, not produced
    extra: list[Row] = field(default_factory=list)              # produced, not in Excel
    outside: list[Row] = field(default_factory=list)            # produced after the last Excel row
    out_of_scope: int = 0                                       # alarm / admin rows (SOC)
    window: tuple[int, int] | None = None


@dataclass
class EvalResult:
    days: list[DayEval]
    fields: tuple[str, ...] = FIELDS

    def totals(self) -> dict:
        n_actual = sum(len(d.pairs) + len(d.missing) for d in self.days)
        n_pred = sum(len(d.pairs) + len(d.extra) for d in self.days)
        n_match = sum(len(d.pairs) for d in self.days)
        exact = sum(1 for d in self.days for p in d.pairs if not p.diffs)
        field_ok = {f: sum(1 for d in self.days for p in d.pairs if f not in p.diffs)
                    for f in self.fields}
        near_time = {f: sum(1 for d in self.days for p in d.pairs
                            if f not in p.diffs or _near(p.diffs[f]))
                     for f in ("D", "E", "F", "G") if f in self.fields}
        recall = n_match / n_actual if n_actual else 0.0
        precision = n_match / n_pred if n_pred else 0.0
        f1 = 2 * recall * precision / (recall + precision) if recall + precision else 0.0
        return {
            "actual_rows": n_actual, "predicted_rows": n_pred, "matched": n_match,
            "recall": recall, "precision": precision, "f1": f1,
            "exact_rows": exact,
            "exact_row_rate": exact / max(n_actual, n_pred) if max(n_actual, n_pred) else 0.0,
            "field_accuracy": {f: (field_ok[f] / n_match if n_match else 0.0) for f in self.fields},
            "time_within_1min": {f: (v / n_match if n_match else 0.0) for f, v in near_time.items()},
            "outside_window": sum(len(d.outside) for d in self.days),
            "out_of_scope": sum(d.out_of_scope for d in self.days),
        }


# ------------------------------------------------------------------ loading
def _seconds(v) -> int | None:
    if v is None or v == "":
        return None
    if isinstance(v, timedelta):
        return int(round(v.total_seconds()))
    if isinstance(v, (datetime, time)):
        return v.hour * 3600 + v.minute * 60 + v.second
    if isinstance(v, (int, float)):
        return int(round(float(v) * 86400)) % 86400
    return None


_MINUS_MINUTE = re.compile(r"=\s*([D-G])\d+\s*-\s*\(\s*1\s*/\s*1440\s*\)")


def _row_times(vals: dict) -> tuple:
    """D..G in seconds. Handles the '=E13-(1/1440)' formulas found in a few sheets."""
    out: dict[str, int | None] = {}
    for col in "GFED":                      # right to left: formulas point to the right
        v = vals.get(col)
        m = _MINUS_MINUTE.fullmatch(v.strip()) if isinstance(v, str) else None
        if m and out.get(m.group(1)) is not None:
            out[col] = out[m.group(1)] - 60
        else:
            out[col] = _seconds(v)
    return tuple(out[c] for c in "DEFG")


def load_actual(path: str | Path, days: list[date] | None = None,
                first_row: int = 10) -> dict[date, tuple[list[ActualRow], int]]:
    """-> {day: (rows of our actions, number of out-of-scope rows)}"""
    wb = openpyxl.load_workbook(path, data_only=False)
    out = {}
    for ws in wb.worksheets:
        if not re.fullmatch(r"\d{4}", ws.title):
            continue
        year = _guess_year(wb)
        d = date(year, int(ws.title[:2]), int(ws.title[2:]))
        if days is not None and d not in days:
            continue
        rows, oos = [], 0
        for i, r in enumerate(ws.iter_rows(min_row=first_row, max_col=12, values_only=True),
                              start=first_row):
            if all(v in (None, "") for v in r[:7]):
                continue
            action = r[2].strip() if isinstance(r[2], str) else r[2]
            if action not in OUR_ACTIONS:
                oos += 1
                continue
            times = _row_times(dict(zip("ABCDEFGHIJKL", r)))
            rows.append(ActualRow(d, i, _s(r[0]), _s(r[1]), action, times, _s(r[10])))
        out[d] = (rows, oos)
    return out


def _guess_year(wb) -> int:
    try:
        v = wb["Daily Totals"]["B2"].value
        if isinstance(v, datetime):
            return v.year
    except KeyError:
        pass
    return datetime.now().year


def _s(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s or None


# ------------------------------------------------------------------ comparison
def norm_name(n: str | None) -> str:
    if not n:
        return ""
    toks = re.sub(r"[^\w\s]", " ", n.lower()).split()
    return " ".join(sorted(toks))


def norm_comment(c: str | None) -> str:
    if not c:
        return ""
    c = c.replace("\\", "").strip()
    m = re.match(r"(?i)\s*metal screening\s+(\S+)\s*:\s*(.*?)\s*-\s*declined to sit\s*$", c)
    if m:
        return f"ms {m.group(1).upper()} : {norm_name(m.group(2))}"
    m = re.match(r"(?i)\s*escort\s+([^\s:]+)\s*:?\s*(.*?)\s*$", c)
    if m:
        return f"escort {m.group(1).lower()} : {norm_name(m.group(2))}"
    return re.sub(r"\s+", " ", c.lower())


def _field(x, f: str):
    if isinstance(x, ActualRow):
        if f in "DEFG":
            return x.times["DEFG".index(f)]
        return {"shift": x.shift, "officer": norm_name(x.officer), "action": x.action,
                "comment": norm_comment(x.comment)}[f]
    d, e, fs, g = x.times_seconds()
    if f in "DEFG":
        return {"D": d, "E": e, "F": fs, "G": g}[f]
    return {"shift": x.shift, "officer": norm_name(x.officer), "action": x.action,
            "comment": norm_comment(x.comment)}[f]


def _near(diff: tuple) -> bool:
    a, b = diff
    return a is not None and b is not None and abs(a - b) <= 60


def _cost(p: Row, a: ActualRow) -> float | None:
    if p.action != a.action:
        return None
    pf, pg = _field(p, "F"), _field(p, "G")
    af, ag = a.times[2], a.times[3]
    if pf is None or af is None:
        return None
    dt = abs(pf - af) / 60 + (abs(pg - ag) / 60 if pg is not None and ag is not None else 5)
    if dt > 30:
        return None
    c = dt
    if _field(p, "officer") != _field(a, "officer"):
        c += 4
    if _field(p, "comment") != _field(a, "comment"):
        c += 4
    return c


def compare_day(day: date, pred: list[Row], actual: list[ActualRow], oos: int,
                fields: tuple[str, ...] = FIELDS) -> DayEval:
    ev = DayEval(day, out_of_scope=oos)
    work = [a for a in actual if a.action != "Assumed Post"]
    if work:
        f0 = min(a.times[2] for a in work if a.times[2] is not None)
        g1 = max((a.times[3] or a.times[2] or 0) for a in work)
        ev.window = (f0, g1)
    cands = []
    for i, p in enumerate(pred):
        for j, a in enumerate(actual):
            c = _cost(p, a)
            if c is not None:
                cands.append((c, i, j))
    cands.sort()
    used_p, used_a = set(), set()
    for c, i, j in cands:
        if i in used_p or j in used_a:
            continue
        used_p.add(i)
        used_a.add(j)
        p, a = pred[i], actual[j]
        diffs = {f: (_field(p, f), _field(a, f)) for f in fields if _field(p, f) != _field(a, f)}
        ev.pairs.append(Pair(p, a, diffs))
    ev.missing = [a for j, a in enumerate(actual) if j not in used_a]
    for i, p in enumerate(pred):
        if i in used_p:
            continue
        start = _field(p, "F")
        if ev.window and start is not None and start > ev.window[1] + 5 * 60:
            ev.outside.append(p)
        else:
            ev.extra.append(p)
    return ev


def evaluate(pred_by_day: dict[date, list[Row]], actual_by_day: dict, days: list[date],
             fields: tuple[str, ...] = FIELDS) -> EvalResult:
    out = []
    for d in days:
        actual, oos = actual_by_day.get(d, ([], 0))
        out.append(compare_day(d, pred_by_day.get(d, []), actual, oos, fields))
    return EvalResult(out, fields)


def describe(x) -> str:
    if isinstance(x, ActualRow):
        d, e, f, g = x.times
        return (f"[שורה {x.sheet_row}] {x.shift} | {x.officer} | {x.action} | "
                f"{fmt_seconds(f)}-{fmt_seconds(g)} | {x.comment or ''}")
    d, e, f, g = x.times_seconds()
    return f"{x.shift} | {x.officer} | {x.action} | {fmt_seconds(f)}-{fmt_seconds(g)} | {x.comment or ''}"


def format_summary(res: EvalResult, title: str = "") -> str:
    t = res.totals()
    lines = [title] if title else []
    lines.append(f"rows in Excel: {t['actual_rows']}  produced: {t['predicted_rows']}  "
                 f"matched: {t['matched']}  (out of scope SOC rows: {t['out_of_scope']}, "
                 f"produced after last Excel row: {t['outside_window']})")
    lines.append(f"recall {t['recall']:.1%}  precision {t['precision']:.1%}  F1 {t['f1']:.1%}  "
                 f"exact rows {t['exact_rows']} ({t['exact_row_rate']:.1%})")
    lines.append("field accuracy on matched rows: " + "  ".join(
        f"{f} {v:.1%}" for f, v in t["field_accuracy"].items()))
    if t["time_within_1min"]:
        lines.append("times within 1 minute: " + "  ".join(
            f"{f} {v:.1%}" for f, v in t["time_within_1min"].items()))
    for d in res.days:
        n_act = len(d.pairs) + len(d.missing)
        exact = sum(1 for p in d.pairs if not p.diffs)
        lines.append(f"  {d.day:%d/%m}: excel {n_act:3d} produced {len(d.pairs) + len(d.extra):3d} "
                     f"matched {len(d.pairs):3d} exact {exact:3d} missing {len(d.missing):3d} "
                     f"extra {len(d.extra):3d} after-window {len(d.outside):3d}")
    return "\n".join(lines)


def write_report(path, res: EvalResult, msgs_by_day: dict | None = None, title: str = "") -> None:
    """Gap report: summary, field differences, rows missing / extra, with the chat messages."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    from .build import sources_text

    head = PatternFill("solid", fgColor="D9E2F3")
    wb = Workbook()
    ws = wb.active
    ws.title = "סיכום"
    ws.sheet_view.rightToLeft = True
    for line in format_summary(res, title).splitlines():
        ws.append([line])
    ws.column_dimensions["A"].width = 140

    def sheet(name, headers, widths):
        sh = wb.create_sheet(name)
        sh.sheet_view.rightToLeft = True
        sh.append(headers)
        for i, w in enumerate(widths, start=1):
            sh.column_dimensions[chr(64 + i)].width = w
            sh.cell(1, i).font = Font(bold=True)
            sh.cell(1, i).fill = head
        sh.freeze_panes = "A2"
        return sh

    def near(day, secs, minutes=10):
        if not msgs_by_day or secs is None:
            return ""
        out = []
        for m in msgs_by_day.get(day, []):
            t = m.ts.hour * 3600 + m.ts.minute * 60
            if -minutes * 60 <= t - secs <= 45 * 60:
                out.append(m)
        return sources_text(out[:8])

    d = sheet("פערים בשדות", ["תאריך", "שדה", "אצלנו", "באקסל", "השורה שלנו", "השורה באקסל", "הודעות מקור"],
              [11, 9, 22, 22, 60, 60, 70])
    for de in res.days:
        for p in de.pairs:
            for f, (ours, theirs) in p.diffs.items():
                if f in "DEFG":
                    ours, theirs = fmt_seconds(ours), fmt_seconds(theirs)
                d.append([de.day.strftime("%d/%m/%Y"), f, ours, theirs, describe(p.pred),
                          describe(p.actual), sources_text(p.pred.sources)])
    m = sheet("חסר אצלנו", ["תאריך", "השורה באקסל", "הודעות בצ'אט באותו זמן"], [11, 80, 90])
    for de in res.days:
        for a in de.missing:
            m.append([de.day.strftime("%d/%m/%Y"), describe(a), near(de.day, a.times[2])])
    x = sheet("עודף אצלנו", ["תאריך", "השורה שלנו", "סימונים", "הודעות מקור"], [11, 80, 40, 80])
    for de in res.days:
        for r in de.extra:
            x.append([de.day.strftime("%d/%m/%Y"), describe(r), "\n".join(r.flags),
                      sources_text(r.sources)])
        for r in de.outside:
            x.append([de.day.strftime("%d/%m/%Y"), describe(r), "אחרי השורה האחרונה באקסל (לא הוזן)",
                      sources_text(r.sources)])
    for sh in wb.worksheets[1:]:
        for row in sh.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)
